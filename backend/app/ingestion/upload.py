"""
upload.py — accept a spreadsheet file, safely.

Honest framing first, because it shapes every decision below: Python cannot
detect malware. There is no library that makes this file "virus-checked". What
this module does instead is remove the conditions under which a malicious
spreadsheet could do anything at all:

  * Only .xlsx and .csv are accepted. Macro-bearing formats (.xlsm, .xlsb, the
    old .xls) are rejected by extension AND by inspecting the file itself, so
    renaming evil.xlsm to safe.xlsx does not get past us.
  * The declared extension must match the actual bytes (magic-number check).
  * An .xlsx is a ZIP archive, so it is inspected before extraction: member
    count, compression ratio and total uncompressed size are checked, which is
    what stops a zip bomb. Members with absolute or traversing paths are
    rejected outright.
  * A workbook containing vbaProject.bin is refused — that is macro payload in
    a file claiming not to have macros.
  * openpyxl runs in read_only + data_only mode. It reads XML; it does not
    execute formulas, macros or external links. With defusedxml installed
    (a hard requirement in requirements.txt) the XML parser also refuses
    entity-expansion attacks.
  * The file is never written to disk. It is parsed from memory and discarded,
    so there is nothing to persist, nothing to re-scan, and nothing that falls
    under data-retention rules.
  * Only the header row and the first few data rows are read, so a 200 MB
    workbook cannot turn into 200 MB of memory.

The residual risk after all of that is an unknown vulnerability in openpyxl or
zlib itself. ClamAV as a sidecar container would add a signature scan for known
malware; it is a reasonable addition later and deliberately not a dependency
now. What it would NOT do is make any of the checks above unnecessary.
"""

from __future__ import annotations

import csv
import io
import zipfile

from .models import ParsedTable, SafetyReport
from .paste import build_table
from .sanitize import MAX_DATA_ROWS

# ---- limits ----------------------------------------------------------------
MAX_FILE_BYTES = 2 * 1024 * 1024        # 2 MB: a header plus a few rows is tiny
MAX_ZIP_MEMBERS = 200                    # a normal .xlsx has well under 50
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200              # 200:1 is far past any honest sheet

ALLOWED_EXTENSIONS = {".xlsx", ".csv"}

# Formats that can carry executable content. Listed explicitly so the rejection
# message can say why, rather than "unsupported file type".
MACRO_CAPABLE_EXTENSIONS = {".xlsm", ".xlsb", ".xls", ".xltm", ".ods"}

_ZIP_MAGIC = b"PK\x03\x04"


class UploadRejected(Exception):
    """The file was refused. The message is safe to show the user."""


def _extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def _check_size(data: bytes) -> None:
    if not data:
        raise UploadRejected("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        limit_mb = MAX_FILE_BYTES // (1024 * 1024)
        raise UploadRejected(
            f"The file is larger than {limit_mb} MB. Only the column headers and "
            f"a few rows are needed — export a small sample, or paste the cells instead."
        )


def _inspect_zip(data: bytes) -> zipfile.ZipFile:
    """Open an .xlsx as a ZIP after checking it is not hostile."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise UploadRejected("This is not a readable .xlsx file.") from exc

    members = archive.infolist()
    if len(members) > MAX_ZIP_MEMBERS:
        raise UploadRejected("This workbook has an unusual internal structure and was not opened.")

    total_uncompressed = 0
    for member in members:
        name = member.filename

        # Path traversal or absolute paths inside the archive.
        if name.startswith("/") or ".." in name.replace("\\", "/").split("/"):
            raise UploadRejected("This workbook contains unsafe internal file paths.")

        # Macro payload inside a file claiming to be macro-free.
        if name.lower().endswith("vbaproject.bin"):
            raise UploadRejected(
                "This workbook contains macros. Save it as .xlsx without macros, "
                "or paste the cells instead."
            )

        total_uncompressed += member.file_size
        if member.compress_size > 0:
            ratio = member.file_size / member.compress_size
            if ratio > MAX_COMPRESSION_RATIO and member.file_size > 1024 * 1024:
                raise UploadRejected("This file expands to an unreasonable size and was not opened.")

    if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
        raise UploadRejected("This file expands to an unreasonable size and was not opened.")

    return archive


def _rows_from_xlsx(data: bytes, max_rows: int) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UploadRejected("Spreadsheet upload is unavailable on this server.") from exc

    archive = _inspect_zip(data)
    archive.close()

    workbook = openpyxl.load_workbook(
        io.BytesIO(data),
        read_only=True,     # streams rows; never builds the whole sheet in memory
        data_only=True,     # cached values, never formulas — nothing to evaluate
        keep_links=False,   # ignore links to other workbooks
    )
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(max_row=max_rows, values_only=True):
            if row is None:
                continue
            cells = ["" if value is None else str(value) for value in row]
            if any(cell.strip() for cell in cells):
                rows.append(cells)
            if len(rows) >= max_rows:
                break
        return rows
    finally:
        workbook.close()


def _rows_from_csv(data: bytes, max_rows: int) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on a Romanian Windows machine often exports cp1250/cp1252.
        for encoding in ("cp1250", "cp1252", "latin-1"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - latin-1 decodes any byte string
            raise UploadRejected("The file's text encoding could not be read.")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows: list[list[str]] = []
    for row in reader:
        if any(cell.strip() for cell in row):
            rows.append(list(row))
        if len(rows) >= max_rows:
            break
    return rows


def parse_upload(filename: str, data: bytes, has_header: bool | None = None) -> ParsedTable:
    """Validate and parse an uploaded spreadsheet into the shared representation.

    Raises UploadRejected with a user-facing message when the file is refused.
    """
    extension = _extension(filename)
    _check_size(data)

    if extension in MACRO_CAPABLE_EXTENSIONS:
        raise UploadRejected(
            f"{extension} files can contain macros and are not accepted. "
            f"Save the sheet as .xlsx or .csv, or paste the cells instead."
        )
    if extension not in ALLOWED_EXTENSIONS:
        raise UploadRejected("Only .xlsx and .csv files are accepted.")

    looks_like_zip = data[:4] == _ZIP_MAGIC

    # The bytes must agree with the name: renaming a macro workbook to .xlsx
    # does not change what is inside it, and a ZIP arriving as .csv is not a CSV.
    if extension == ".xlsx" and not looks_like_zip:
        raise UploadRejected("This file is not a real .xlsx workbook.")
    if extension == ".csv" and looks_like_zip:
        raise UploadRejected("This file is named .csv but is actually a compressed workbook.")

    # One extra row is read so the report can tell the user their table was
    # longer than the sample we used.
    read_limit = MAX_DATA_ROWS + 2

    if extension == ".xlsx":
        grid = _rows_from_xlsx(data, read_limit)
        source = "xlsx"
    else:
        grid = _rows_from_csv(data, read_limit)
        source = "csv"

    if not grid:
        raise UploadRejected("No rows could be read from this file.")

    report = SafetyReport(source=source, original_chars=len(data))
    table = build_table(grid, report, has_header=has_header)

    # We stopped reading early, so `rows_seen` is what we read, not what the
    # file holds. Say that plainly rather than reporting a count we did not
    # measure — a wrong number is worse than an honest "at least".
    if report.truncated_rows:
        report.warnings = [
            w for w in report.warnings if not w.startswith("using the first")
        ]
        report.warnings.append(
            f"using the first {MAX_DATA_ROWS} rows of the file — that is all the "
            f"engine needs to identify your columns"
        )

    return table
